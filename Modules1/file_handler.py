# modules/file_handler.py

import docx
import PyPDF2
from fpdf import FPDF
import io
import os
def read_file(file):
    extension = file.name.split('.')[-1].lower()

    if extension == 'pdf':
        return read_pdf(file)
    elif extension == 'docx':
        return read_docx(file)
    elif extension == 'txt':
        return file.read().decode('utf-8')
    else:
        return "❌ Unsupported file format"

def read_pdf(file):
    reader = PyPDF2.PdfReader(file)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text

def read_docx(file):
    doc = docx.Document(file)
    return "\n".join([para.text for para in doc.paragraphs])

def save_summary_as_txt(text):
    return io.BytesIO(text.encode("utf-8"))

# def save_summary_as_pdf(text):
#     pdf = FPDF()
#     pdf.add_page()
#     pdf.set_auto_page_break(auto=True, margin=15)
#     pdf.set_font("Arial", size=12)
#     for line in text.split('\n'):
#         pdf.cell(200, 10, txt=line, ln=True)
#     output = io.BytesIO()
#     pdf.output(output)
#     output.seek(0)
#     return output
 


# def save_summary_as_pdf(text):
#     pdf = FPDF()
#     pdf.add_page()

#     # Add a Unicode font (e.g., DejaVuSans)
#     font_path = os.path.join(os.path.dirname(__file__), "DejaVuSans.ttf")
#     if not os.path.exists(font_path):
#         # Auto-download the font if not present
#         import urllib.request
#         url = "https://github.com/dejavu-fonts/dejavu-fonts/raw/version_2_37/ttf/DejaVuSans.ttf"
#         urllib.request.urlretrieve(url, font_path)

#     pdf.add_font("DejaVu", "", font_path, uni=True)
#     pdf.set_font("DejaVu", size=12)

#     pdf.set_auto_page_break(auto=True, margin=15)

#     for line in text.split('\n'):
#         pdf.multi_cell(0, 10, line)

#     output = io.BytesIO()
#     pdf.output(output)
#     output.seek(0)
#     return output

from fpdf import FPDF

def save_summary_as_pdf(summary_text):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)  # Default font
    
    for line in summary_text.split("\n"):
        pdf.multi_cell(0, 10, line)
    
    output_path = "summary.pdf"
    pdf.output(output_path)
    return output_path

